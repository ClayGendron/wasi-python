"""lxml probe: run natively and in the guest; DATA is a folder (read-only in the guest) holding ext.txt and doc.xml; OUT a writable folder or ''."""
import io, sys, os
DATA = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else ""

def show(label, f):
    try:
        print(label, "->", f())
    except BaseException as e:
        print(label, "-> RAISES", type(e).__name__, str(e)[:160].replace(DATA, "<DATA>"))

from lxml import etree, objectify
show("libxml compiled", lambda: etree.LIBXML_COMPILED_VERSION)
show("libxml runtime", lambda: etree.LIBXML_VERSION)
show("libxslt", lambda: etree.LIBXSLT_VERSION)
show("features", lambda: sorted(etree.LIBXML_FEATURES) if hasattr(etree, "LIBXML_FEATURES") else None)
show("has iconv/zlib", lambda: (etree.LIBXML_FEATURES & {"iconv", "zlib", "thread", "http"}) if hasattr(etree, "LIBXML_FEATURES") else None)

doc = etree.fromstring(b"<r xmlns:a='urn:a'><i n='1'>x</i><i n='2'>y<a:b/></i><!-- c --><?pi d?></r>")
show("xpath count", lambda: doc.xpath("count(//i)"))
show("xpath ns", lambda: [e.tag for e in doc.xpath("//a:b", namespaces={"a": "urn:a"})])
show("xpath str", lambda: doc.xpath("string(//i[@n='2'])"))
show("xpath fn", lambda: doc.xpath("concat(name(/*), '-', sum(//i/@n))"))
show("exslt regex", lambda: doc.xpath("//i[re:test(., '^y')]/@n", namespaces={"re": "http://exslt.org/regular-expressions"}))
show("tostring", lambda: etree.tostring(doc))
show("c14n", lambda: etree.tostring(doc, method="c14n"))
show("pretty", lambda: etree.tostring(etree.fromstring(b"<a><b>1</b><c/></a>"), pretty_print=True))
xslt = etree.XSLT(etree.fromstring(b"""<xsl:stylesheet version='1.0' xmlns:xsl='http://www.w3.org/1999/XSL/Transform'
 xmlns:str='http://exslt.org/strings' xmlns:math='http://exslt.org/math' extension-element-prefixes='str math'>
<xsl:output method='xml' indent='no'/>
<xsl:template match='/'><out total='{sum(//i/@n)}' max='{math:max(//i/@n)}'><xsl:for-each select='//i'><xsl:sort select='@n' data-type='number' order='descending'/>
<v><xsl:value-of select='str:padding(3, "*")'/><xsl:value-of select='.'/></v></xsl:for-each></out></xsl:template></xsl:stylesheet>"""))
show("xslt", lambda: str(xslt(doc)))
show("xslt html", lambda: str(etree.XSLT(etree.XML(b"<xsl:stylesheet version='1.0' xmlns:xsl='http://www.w3.org/1999/XSL/Transform'><xsl:output method='html'/><xsl:template match='/'><html><head><title>t</title></head><body><p><xsl:value-of select='count(//i)'/></p></body></html></xsl:template></xsl:stylesheet>"))(doc)))
show("objectify", lambda: (lambda o: (o.a + 1, o.b.text, type(o.c).__name__))(objectify.fromstring(b"<r><a>41</a><b>s</b><c>1.5</c></r>")))
show("html parse", lambda: etree.tostring(etree.HTML(b"<p>a<b>b<p>c"), method="html"))
from lxml import html as lh
show("lxml.html", lambda: [a.get("href") for a in lh.fromstring("<div><a href='/x'>1</a><a href='/y'>2</a></div>").iter("a")])
show("lxml.html text", lambda: lh.fromstring("<div>a <b>b</b> c</div>").text_content())
show("html tostring meta", lambda: lh.tostring(lh.document_fromstring("<html><head><meta charset='utf-8'><title>x</title></head><body>é</body></html>")))
show("cssselect absent", lambda: __import__("lxml.cssselect"))
# encodings
for enc, txt in (("windows-1252", "café €"), ("iso-8859-2", "Łódź"), ("koi8-r", "привет"), ("shift_jis", "日本"), ("euc-jp", "日本"), ("gb2312", "中文"), ("big5", "中文"), ("utf-16", "xé")):
    data = ('<?xml version="1.0" encoding="%s"?><r>%s</r>' % (enc, txt)).encode(enc)
    show("parse " + enc, lambda: etree.fromstring(data).text == txt)
    show("write " + enc, lambda: etree.tostring(etree.fromstring("<r>%s</r>" % txt), encoding=enc) == ("<?xml version='1.0' encoding='%s'?>\n<r>%s</r>" % (enc, txt)).encode(enc))
# safety defaults
xxe = b"<!DOCTYPE r [<!ENTITY e SYSTEM 'file://" + DATA.encode() + b"/ext.txt'>]><r>&e;</r>"
show("xxe default", lambda: etree.tostring(etree.fromstring(xxe)).replace(DATA.encode(), b"<DATA>"))
show("xxe resolve", lambda: etree.tostring(etree.fromstring(xxe, etree.XMLParser(resolve_entities=True))).replace(DATA.encode(), b"<DATA>"))
show("xxe net", lambda: etree.tostring(etree.fromstring(b"<!DOCTYPE r [<!ENTITY e SYSTEM 'http://127.0.0.1:9/x'>]><r>&e;</r>", etree.XMLParser(resolve_entities=True, no_network=False))))
show("billion laughs", lambda: len(etree.tostring(etree.fromstring(b"<!DOCTYPE r [<!ENTITY a 'aaaaaaaaaa'><!ENTITY b '&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;'><!ENTITY c '&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;'><!ENTITY d '&c;&c;&c;&c;&c;&c;&c;&c;&c;&c;'><!ENTITY e '&d;&d;&d;&d;&d;&d;&d;&d;&d;&d;'><!ENTITY f '&e;&e;&e;&e;&e;&e;&e;&e;&e;&e;'><!ENTITY g '&f;&f;&f;&f;&f;&f;&f;&f;&f;&f;'><!ENTITY h '&g;&g;&g;&g;&g;&g;&g;&g;&g;&g;'>]><r>&h;</r>", etree.XMLParser(resolve_entities=True)))))
# files by path (libxml2's own I/O: the dup() path)
show("parse path", lambda: etree.parse(os.path.join(DATA, "doc.xml")).getroot().get("k"))
show("parse path twice", lambda: [etree.parse(os.path.join(DATA, "doc.xml")).getroot().tag for _ in range(50)][-1])
show("parse missing", lambda: etree.parse(os.path.join(DATA, "nope.xml")))
show("iterparse path", lambda: [(ev, el.tag) for ev, el in etree.iterparse(os.path.join(DATA, "doc.xml"))])
show("xinclude", lambda: (lambda t: (t.xinclude(), etree.tostring(t))[1])(etree.parse(os.path.join(DATA, "inc.xml"))))
show("xslt document()", lambda: str(etree.XSLT(etree.XML(("<xsl:stylesheet version='1.0' xmlns:xsl='http://www.w3.org/1999/XSL/Transform'><xsl:template match='/'><o><xsl:value-of select=\"document('file://%s/doc.xml')/*/@k\"/></o></xsl:template></xsl:stylesheet>" % DATA).encode()))(doc)))
show("dtd validate", lambda: etree.DTD(os.path.join(DATA, "d.dtd")).validate(etree.parse(os.path.join(DATA, "doc.xml"))))
show("xsd", lambda: etree.XMLSchema(etree.XML(b"<xs:schema xmlns:xs='http://www.w3.org/2001/XMLSchema'><xs:element name='r' type='xs:int'/></xs:schema>")).validate(etree.XML(b"<r>5</r>")))
show("relaxng", lambda: etree.RelaxNG(etree.XML(b"<element name='r' xmlns='http://relaxng.org/ns/structure/1.0'><text/></element>")).validate(etree.XML(b"<r>t</r>")))
show("schematron", lambda: etree.Schematron(etree.XML(b"<schema xmlns='http://www.ascc.net/xml/schematron'><pattern name='p'><rule context='r'><assert test='@k'>k</assert></rule></pattern></schema>")).validate(etree.XML(b"<r k='1'/>")))
show("isoschematron", lambda: __import__("lxml.isoschematron").isoschematron.Schematron)
if OUT:
    def w():
        etree.ElementTree(etree.XML(b"<w>1</w>")).write(os.path.join(OUT, "w.xml"))
        return open(os.path.join(OUT, "w.xml"), "rb").read()
    show("write path", w)
    def wz():
        etree.ElementTree(etree.XML(b"<w>1</w>")).write(os.path.join(OUT, "w.xml.gz"), compression=6)
        import gzip
        return gzip.open(os.path.join(OUT, "w.xml.gz")).read()
    show("write gz", wz)
    show("parse gz decompress", lambda: etree.parse(os.path.join(OUT, "w.xml.gz"), etree.XMLParser(decompress=True)).getroot().text)
show("write stdout -", lambda: etree.ElementTree(etree.XML(b"<o>-</o>")).write("-") )
print()
import pandas as pd
html = "<table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>x</td></tr><tr><td>2</td><td>y</td></tr></table>"
show("read_html default", lambda: pd.read_html(io.StringIO(html))[0].to_dict("list"))
show("read_html lxml", lambda: pd.read_html(io.StringIO(html), flavor="lxml")[0].to_dict("list"))
xml = "<data><row><a>1</a><b>x</b></row><row><a>2</a><b>y</b></row></data>"
show("read_xml lxml", lambda: pd.read_xml(io.StringIO(xml)).to_dict("list"))
show("read_xml xpath", lambda: pd.read_xml(io.StringIO(xml), xpath="//row[a>1]").to_dict("list"))
show("read_xml stylesheet", lambda: pd.read_xml(io.StringIO(xml), stylesheet=io.StringIO("<xsl:stylesheet version='1.0' xmlns:xsl='http://www.w3.org/1999/XSL/Transform'><xsl:template match='/'><data><xsl:for-each select='//row'><r><s><xsl:value-of select='a*10'/></s></r></xsl:for-each></data></xsl:template></xsl:stylesheet>")).to_dict("list"))
show("to_xml lxml", lambda: pd.DataFrame({"a": [1, 2], "b": ["x", "y"]}).to_xml(parser="lxml"))
show("to_xml etree", lambda: pd.DataFrame({"a": [1, 2], "b": ["x", "y"]}).to_xml(parser="etree"))
print()
import docx
show("docx Document()", lambda: len(docx.Document().paragraphs))
