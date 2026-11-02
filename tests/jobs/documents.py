"""Golden job: a Word document and a PowerPoint deck, each from the package's default template.

Usage: documents.py <data dir> <out dir>. python-docx starts from `Document()`, adds a heading,
bold and italic runs, a 3x3 table, a picture (a PNG Pillow writes in memory), a section's own
header and footer, and a comment. python-pptx starts from `Presentation()`, adds a title slide, a
bullet slide, a table, a column chart (xlsxwriter writes its workbook), a picture and a notes
slide. Each file is saved to the output folder, reopened, and shown as what it reads back, then
removed: the archive's bytes are never shown, since member times, core-property dates and the
comment's date come from the clock.
"""

import io
import os
import sys
import zipfile
from importlib.metadata import version

import docx
import pptx
from docx.shared import Inches as DocxInches
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

data_dir, out_dir = sys.argv[1], sys.argv[2]
names = ("python-docx", "python-pptx", "lxml", "pillow", "xlsxwriter")
print("versions", " ".join(f"{d}={version(d)}" for d in names))


def png() -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def members(path: str) -> None:
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            print("  member", n)


# 1. python-docx: a document from the default template, saved, reopened, read back
d = docx.Document()
d.add_heading("Quarterly report", level=1)
p = d.add_paragraph("Sales were ")
p.add_run("up").bold = True
p.add_run(" in the ")
p.add_run("north").italic = True
p.add_run(".")
table = d.add_table(rows=3, cols=3)
for r, row in enumerate(table.rows):
    for c, cell in enumerate(row.cells):
        cell.text = f"r{r}c{c}"
d.add_picture(png(), width=DocxInches(1))
section = d.sections[0]
section.different_first_page_header_footer = False
section.header.is_linked_to_previous = False
section.header.paragraphs[0].text = "the header"
section.footer.is_linked_to_previous = False
section.footer.paragraphs[0].text = "the footer"
d.add_comment(p.runs[1], text="check this figure", author="Reviewer", initials="RV")
path = os.path.join(out_dir, "report.docx")
d.save(path)

print("docx members:")
members(path)
back = docx.Document(path)
for para in back.paragraphs:
    runs = [(r.text, r.bold, r.italic) for r in para.runs]
    print("  paragraph", para.style.name, repr(para.text), runs)
for t in back.tables:
    print("  table", [[cell.text for cell in row.cells] for row in t.rows])
for shape in back.inline_shapes:
    print("  inline shape", shape.width, shape.height)
for part in back.part.package.image_parts:
    print("  image", part.content_type, part.image.px_width, part.image.px_height)
s = back.sections[0]
print("  header", s.header.is_linked_to_previous, [q.text for q in s.header.paragraphs])
print("  footer", s.footer.is_linked_to_previous, [q.text for q in s.footer.paragraphs])
for comment in back.comments:
    print("  comment", comment.author, comment.initials, repr(comment.text))
print("  styles", len(back.styles), "default font", back.styles["Normal"].font.name)
print("  settings odd and even pages", back.settings.odd_and_even_pages_header_footer)
os.remove(path)

# 2. python-pptx: a deck from the default template, saved, reopened, read back
prs = pptx.Presentation()
title = prs.slides.add_slide(prs.slide_layouts[0])
title.shapes.title.text = "Quarterly report"
title.placeholders[1].text = "Sales by region"
bullets = prs.slides.add_slide(prs.slide_layouts[1])
bullets.shapes.title.text = "Highlights"
frame = bullets.placeholders[1].text_frame
frame.text = "North up"
for text, level in (("South flat", 0), ("East new", 1)):
    para = frame.add_paragraph()
    para.text = text
    para.level = level
blank = prs.slides.add_slide(prs.slide_layouts[6])
shape = blank.shapes.add_table(2, 3, Inches(0.5), Inches(0.5), Inches(6), Inches(1))
for r, row in enumerate(shape.table.rows):
    for c, cell in enumerate(row.cells):
        cell.text = f"r{r}c{c}"
chart_data = CategoryChartData()
chart_data.categories = ["north", "south", "east"]
chart_data.add_series("units", (3, 5, 2))
chart_data.add_series("returns", (1, 0, 1))
blank.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2), Inches(6), Inches(4), chart_data
)
blank.shapes.add_picture(png(), Inches(7), Inches(0.5), width=Inches(1))
blank.notes_slide.notes_text_frame.text = "Say the totals."
path = os.path.join(out_dir, "report.pptx")
prs.save(path)

print("pptx members:")
members(path)
back = pptx.Presentation(path)
print("  slide size", back.slide_width, back.slide_height)
for n, slide in enumerate(back.slides):
    print("  slide", n, slide.slide_layout.name, "notes", slide.has_notes_slide)
    for sh in slide.shapes:
        print("    shape", sh.shape_type, sh.name, sh.left, sh.top, sh.width, sh.height)
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                print("      paragraph", para.level, repr(para.text))
        if sh.has_table:
            print("      table", [[cell.text for cell in row.cells] for row in sh.table.rows])
        if sh.has_chart:
            chart = sh.chart
            print("      chart", chart.chart_type, list(chart.plots[0].categories))
            for series in chart.series:
                print("      series", series.name, list(series.values))
        if sh.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print("      image", sh.image.content_type, sh.image.size)
    if slide.has_notes_slide:
        print("    notes", repr(slide.notes_slide.notes_text_frame.text))
os.remove(path)
