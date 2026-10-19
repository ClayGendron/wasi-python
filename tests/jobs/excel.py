"""Golden job: Excel files, read and written (spec 039, Decisions 9 and 11).

Usage: excel.py <data dir> <out dir>. xlsxwriter writes a workbook in memory with a fixed creation
date, so its parts are the same bytes every time; each part is shown as its SHA-256, never the
deflated container (zlib differs between the oracle and the guest). openpyxl stamps the save time
into its output, so what it writes is read back and shown as values. pandas reads with openpyxl and
writes through xlsxwriter in memory, the one writer that needs no writable folder; openpyxl writes
after `tempfile.tempdir` names a writable folder (the README's pattern). xlrd refuses an .xlsx with
its own error.
"""

import datetime
import hashlib
import io
import sys
import tempfile
import zipfile
from importlib.metadata import version

import openpyxl
import pandas as pd
import xlrd
import xlsxwriter

data_dir, out_dir = sys.argv[1], sys.argv[2]
names = ("openpyxl", "et-xmlfile", "xlsxwriter", "xlrd", "pandas")
print("versions", " ".join(f"{d}={version(d)}" for d in names))


def parts(blob):
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        for n in sorted(z.namelist()):
            print("  part", n, hashlib.sha256(z.read(n)).hexdigest()[:16])


# 1. xlsxwriter, in memory: nothing on disk
b = io.BytesIO()
wb = xlsxwriter.Workbook(b, {"in_memory": True})
wb.set_properties({"created": datetime.datetime(2026, 1, 1)})
ws = wb.add_worksheet("sales")
ws.write_row(0, 0, ["region", "units", "price", "day"])
rows = [
    ("north", 3, 2.5, datetime.date(2026, 3, 1)),
    ("south", 5, 1.25, datetime.date(2026, 3, 2)),
]
datefmt = wb.add_format({"num_format": "yyyy-mm-dd"})
for i, (region, units, price, day) in enumerate(rows, 1):
    ws.write_row(i, 0, [region, units, price])
    ws.write_datetime(i, 3, datetime.datetime.combine(day, datetime.time()), datefmt)
ws.write_formula(3, 1, "=SUM(B2:B3)", None, 8)
wb.close()
blob = b.getvalue()
print("xlsxwriter parts:")
parts(blob)

# 2. openpyxl reads it back: the rows, the formula, and its cached value
book = openpyxl.load_workbook(io.BytesIO(blob))
sheet = book["sales"]
for row in sheet.iter_rows(values_only=True):
    print("  row", row)
print("  formula", sheet["B4"].value)
cached = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
print("  cached value", cached["sales"]["B4"].value)

# 3. pandas reads with openpyxl, writes a groupby with xlsxwriter in memory, and reads it again
df = pd.read_excel(io.BytesIO(blob), sheet_name="sales", nrows=2)
print(df.to_csv(index=False), end="")
out = io.BytesIO()
options = {"options": {"in_memory": True}}
with pd.ExcelWriter(out, engine="xlsxwriter", engine_kwargs=options) as w:
    totals = df.groupby("region", as_index=False)["units"].sum()
    totals.to_excel(w, sheet_name="totals", index=False)
print(pd.read_excel(io.BytesIO(out.getvalue()), sheet_name="totals").to_csv(index=False), end="")

# 4. openpyxl writes each sheet through tempfile: it is pointed at the writable data folder
tempfile.tempdir = data_dir
wb2 = openpyxl.Workbook()
wb2.active.title = "t"
wb2.active.append(["a", 1])
wb2.active.append(["b", 2.5])
saved = io.BytesIO()
wb2.save(saved)
print("openpyxl round trip", list(openpyxl.load_workbook(io.BytesIO(saved.getvalue()))["t"].values))

# 5. sales.csv through xlsxwriter (in memory) and back through openpyxl
sales = pd.read_csv(f"{data_dir}/sales.csv")
buf = io.BytesIO()
sales.to_excel(buf, engine="xlsxwriter", engine_kwargs=options, index=False)
back = pd.read_excel(io.BytesIO(buf.getvalue()))
by_region = back.groupby("region")["sales"].sum().round(2).to_dict()
print("sales round trip", back.equals(sales), by_region)

# 6. xlrd reads only .xls, and refuses an .xlsx with its own error
try:
    xlrd.open_workbook(file_contents=blob)
except xlrd.biffh.XLRDError as e:
    print("xlrd:", e)
